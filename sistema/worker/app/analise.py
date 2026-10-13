"""Grava a análise dos agentes de IA (Markdown + JSON) e gera os modelos de declaração (.docx)."""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .documentos import SUB_ANALISE, SUB_DECLARACOES
from .naming import limpar_nome, pasta_processo

CAMPOS_PRAZO = [
    ("validade_proposta", "Validade da proposta"),
    ("validade_produtos", "Validade mínima dos produtos/itens"),
    ("prazo_entrega", "Prazo de entrega após empenho/OF"),
    ("prazo_pagamento", "Prazo de pagamento"),
    ("vigencia_ata_contrato", "Vigência da ata/contrato"),
    ("local_entrega", "Local de entrega"),
]
CAMPOS_CUSTO = [
    ("garantia_proposta", "Garantia de proposta"),
    ("garantia_contratual", "Garantia contratual"),
    ("pagamento_antecipado", "Pagamento antecipado / taxas"),
    ("amostras", "Amostras / prova de conceito"),
    ("visita_tecnica", "Visita técnica"),
    ("exclusividade_me_epp", "Exclusividade / cota ME-EPP"),
]

DADOS_EMPRESA = {
    "razao": "[RAZÃO SOCIAL DA EMPRESA]",
    "cnpj": "[CNPJ]",
    "endereco": "[ENDEREÇO COMPLETO]",
    "representante": "[NOME DO REPRESENTANTE LEGAL]",
    "cpf": "[CPF DO REPRESENTANTE]",
    "cargo": "[CARGO]",
}


def _v(campo) -> tuple[str, str]:
    if isinstance(campo, dict):
        return str(campo.get("valor") or "NÃO HÁ INFORMAÇÃO SUFICIENTE"), str(campo.get("fonte") or "—")
    if campo in (None, ""):
        return "NÃO HÁ INFORMAÇÃO SUFICIENTE", "—"
    return str(campo), "—"


def resumo_para_relatorio(a: dict) -> dict:
    """Campos achatados usados na aba de detalhes do Excel."""
    p, g = a.get("prazos", {}) or {}, a.get("garantias_e_custos", {}) or {}
    out = {k: _v(p.get(k))[0] for k, _ in CAMPOS_PRAZO}
    out.update({k: _v(g.get(k))[0] for k, _ in CAMPOS_CUSTO})
    out["relevancia"] = a.get("relevancia")
    out["resumo_ia"] = a.get("resumo")
    out["qtd_declaracoes"] = len(a.get("declaracoes") or [])
    out["riscos"] = "; ".join(str(r) for r in (a.get("riscos") or []))[:1500]
    out["esclarecimentos"] = "; ".join(str(r) for r in (a.get("pontos_esclarecimento_impugnacao") or []))[:1500]
    out["confianca"] = (a.get("auditoria") or {}).get("confianca")
    return out


def _brl(v) -> str:
    try:
        v = float(v)
    except (TypeError, ValueError):
        v = 0.0
    if not v:
        return "não informado (sigiloso ou não divulgado no PNCP)"
    return "R$ " + f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _data_br(v) -> str:
    try:
        return datetime.fromisoformat(str(v)[:19]).strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return str(v or "—")


_NUL_HEX = re.compile(r"\x00([0-9a-fA-F]{2})")


def _reparar(v):
    """Corrige acentos corrompidos vindos de algumas respostas de IA ("Declara\\x00e7\\x00e3o" -> "Declaração")."""
    if isinstance(v, str):
        return _NUL_HEX.sub(lambda m: chr(int(m.group(1), 16)), v).replace("\x00", "")
    if isinstance(v, list):
        return [_reparar(x) for x in v]
    if isinstance(v, dict):
        return {k: _reparar(x) for k, x in v.items()}
    return v


def salvar_analise(pasta: Path, rec: dict, analise: dict, modelos: dict | None = None) -> dict:
    analise = _reparar(analise)
    info = pasta_processo(rec)
    rec = {**rec, "numero_processo": rec.get("numero_processo") or info["numero"], "portal": rec.get("portal") or info["portal"]}
    dados = {**DADOS_EMPRESA, **(modelos or {})}
    pa = pasta / SUB_ANALISE
    pd_ = pasta / SUB_DECLARACOES
    pa.mkdir(parents=True, exist_ok=True)
    pd_.mkdir(parents=True, exist_ok=True)
    (pa / "analise_ia.json").write_text(json.dumps(analise, ensure_ascii=False, indent=2), encoding="utf-8")

    md = _markdown(rec, analise)
    (pasta / "00_RESUMO_DO_PROCESSO.md").write_text(md, encoding="utf-8")

    for antigo in pd_.glob("*.docx"):  # reprocessamento: remove modelos da execução anterior (evita duplicados)
        antigo.unlink()
    gerados = []
    decls = analise.get("declaracoes") or []
    for i, d in enumerate(decls, 1):
        nome = limpar_nome(d.get("nome") or f"Declaração {i}", 80).replace(" ", "_")
        arq = pd_ / f"{i:02d}_{nome}.docx"
        _docx_declaracao(arq, rec, d, dados)
        gerados.append(arq.name)
    (pd_ / "00_LISTA_DECLARACOES_EXIGIDAS.md").write_text(_lista_decl(rec, decls), encoding="utf-8")
    return {"resumo": str(pasta / "00_RESUMO_DO_PROCESSO.md"), "declaracoes": gerados}


def _markdown(rec: dict, a: dict) -> str:
    L = [
        f"# {rec.get('orgao')} — {rec.get('numero_processo')}",
        "",
        f"- **Objeto:** {rec.get('objetoCompra')}",
        f"- **Encerramento das propostas:** {_data_br(rec.get('dataEncerramentoProposta'))}",
        f"- **Modalidade:** {rec.get('modalidadeNome')} | **Portal:** {rec.get('portal')} | **SRP:** {'Sim' if rec.get('srp') else 'Não'}",
        f"- **Valor estimado:** {_brl(rec.get('valorTotalEstimado'))}",
        f"- **PNCP:** {rec.get('link_pncp')}",
        f"- **Portal de disputa:** {rec.get('linkSistemaOrigem') or '—'}",
        f"- **Relevância (IA):** {a.get('relevancia')} — {a.get('justificativa_relevancia', '')}",
        "",
        "## Resumo",
        str(a.get("resumo") or "—"),
        "",
        "## Prazos e condições",
        "| Campo | Valor | Fonte (arquivo · página · cláusula) |",
        "|---|---|---|",
    ]
    for k, rot in CAMPOS_PRAZO:
        v, f = _v((a.get("prazos") or {}).get(k))
        L.append(f"| {rot} | {v} | {f} |")
    for k, rot in CAMPOS_CUSTO:
        v, f = _v((a.get("garantias_e_custos") or {}).get(k))
        L.append(f"| {rot} | {v} | {f} |")

    def lista(titulo, itens, fmt):
        L.extend(["", f"## {titulo}"])
        if not itens:
            L.append("- NÃO HÁ INFORMAÇÃO SUFICIENTE")
        for x in itens or []:
            L.append("- " + fmt(x))

    fmt_dict = lambda x: (f"{x.get('documento') or x.get('exigencia') or x.get('descricao')} — _{x.get('fonte', '—')}_" if isinstance(x, dict) else str(x))
    lista("Habilitação exigida", a.get("habilitacao"), fmt_dict)
    lista("Exigências técnicas / certificações", a.get("exigencias_tecnicas"), fmt_dict)
    lista("Declarações exigidas (ver 02_MODELOS_DECLARACOES)", a.get("declaracoes"),
          lambda d: f"**{d.get('nome')}** — momento: {d.get('momento', '—')} — _{d.get('fonte', '—')}_")
    lista("Itens aderentes ao portfólio", a.get("itens_aderentes"),
          lambda x: f"Item {x.get('item')}: {x.get('descricao')} — {x.get('comentario', '')}" if isinstance(x, dict) else str(x))
    lista("Riscos", a.get("riscos"), str)
    lista("Pontos para esclarecimento / impugnação", a.get("pontos_esclarecimento_impugnacao"), str)
    lista("Lacunas de informação", a.get("lacunas"), str)
    cz = a.get("cruzamento")
    if cz is not None:
        L.extend(["", "## Cruzamento com as tabelas de preço (ver 04_PROPOSTA_PREVIA)"])
        for av in cz.get("avisos") or []:
            L.append(f"- ⚠ {av}")
        nomes = {k: (v or {}).get("nome") or k for k, v in (cz.get("catalogos") or {}).items()}
        for emp, nome in nomes.items():
            vs = [((it.get("ofertas") or {}).get(emp) or {}).get("veredito") for it in cz.get("itens") or []]
            if any(vs):
                cont = {k: vs.count(k) for k in ("ATENDE", "ATENDE_PARCIAL", "VERIFICAR", "NAO_ATENDE") if vs.count(k)}
                L.append(f"- **{nome}:** " + ", ".join(f"{k.replace('_', ' ').lower()}: {v}" for k, v in cont.items()))
        for it in (cz.get("itens") or [])[:60]:
            ofs = []
            for emp, nome in nomes.items():
                o = (it.get("ofertas") or {}).get(emp)
                if o and o.get("codigo"):
                    ofs.append(f"{nome} cód. {o['codigo']} ({o.get('marca')}) — {o.get('veredito')}")
            L.append(f"- Item {it.get('numero')}: {str(it.get('descricao'))[:90]} → " + ("; ".join(ofs) or "sem correspondência"))
    aud = a.get("auditoria") or {}
    if aud:
        lista(f"Auditoria cruzada (confiança: {aud.get('confianca', '—')})", aud.get("divergencias"), str)
        rev = aud.get("revisor_chatgpt")
        if rev:
            conc = {True: "concorda com a análise", False: "NÃO concorda — ver observações"}.get(rev.get("concorda"), "—")
            lista(f"Revisão independente · ChatGPT ({conc})", rev.get("observacoes"), str)
        if aud.get("avisos"):
            lista("Avisos do processamento", aud.get("avisos"), str)
    L += ["", f"_Gerado automaticamente em {datetime.now():%d/%m/%Y %H:%M}. Conferir sempre no edital antes de enviar a proposta._"]
    return "\n".join(L)


def _lista_decl(rec: dict, decls: list[dict]) -> str:
    L = [f"# Declarações exigidas — {rec.get('orgao')} — {rec.get('numero_processo')}", "",
         "| # | Declaração | Momento | Obrigatória | Onde está no edital | Modelo no edital? |", "|---|---|---|---|---|---|"]
    for i, d in enumerate(decls, 1):
        L.append(f"| {i} | {d.get('nome')} | {d.get('momento', '—')} | {'Sim' if d.get('obrigatoria', True) else 'Não'} | "
                 f"{d.get('fonte', '—')} | {'Sim' if d.get('texto_modelo_edital') else 'Não — modelo sugerido'} |")
    if not decls:
        L.append("| — | NÃO HÁ INFORMAÇÃO SUFICIENTE | | | | |")
    return "\n".join(L)


def _docx_declaracao(arq: Path, rec: dict, d: dict, emp: dict) -> None:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run((d.get("nome") or "DECLARAÇÃO").upper())
    r.bold = True
    r.font.size = Pt(13)
    doc.add_paragraph(f"Ao(À) {rec.get('orgao')}\nRef.: {rec.get('modalidadeNome')} nº {rec.get('numeroCompra')}/{rec.get('anoCompra')} — Processo {rec.get('processo')}")
    corpo = d.get("texto_modelo_edital") or d.get("modelo_sugerido") or ""
    if not corpo:
        corpo = (f"{emp['razao']}, inscrita no CNPJ sob o nº {emp['cnpj']}, com sede em {emp['endereco']}, por intermédio de seu representante legal, "
                 f"{emp['representante']}, CPF nº {emp['cpf']}, DECLARA, para os fins do edital em referência, {d.get('conteudo') or '[CONTEÚDO EXIGIDO]'}.")
    for k, v in {"[RAZÃO SOCIAL]": emp["razao"], "[CNPJ]": emp["cnpj"], "[ENDEREÇO]": emp["endereco"],
                 "[REPRESENTANTE]": emp["representante"], "[CPF]": emp["cpf"]}.items():
        corpo = corpo.replace(k, v)
    para = doc.add_paragraph(corpo)
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    doc.add_paragraph(f"\n{emp.get('cidade') or '[LOCAL]'}, [DATA].\n\n\n______________________________________\n"
                      f"{emp['representante']}\n{emp['cargo']} — {emp['razao']}")
    nota = doc.add_paragraph()
    rn = nota.add_run(f"\nOrigem: {d.get('fonte', '—')} | Fundamento: {d.get('fundamento', '—')} | "
                      f"{'Transcrito do modelo do edital' if d.get('texto_modelo_edital') else 'Modelo SUGERIDO pela IA — conferir com o edital'}")
    rn.italic = True
    rn.font.size = Pt(8)
    doc.save(str(arq))
