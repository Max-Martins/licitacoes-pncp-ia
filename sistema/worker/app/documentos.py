"""Download, organização, hash e extração de texto dos documentos de cada processo."""
from __future__ import annotations

import hashlib
import io
import json
import logging
import re
import time
import zipfile
from datetime import datetime
from pathlib import Path

import httpx

from .naming import limpar_nome, sem_acento

log = logging.getLogger("documentos")

# Duas rotas oficiais para os mesmos dados; tenta a pública primeiro e cai para a do portal.
PNCP_BASES = ["https://pncp.gov.br/pncp-api/v1", "https://pncp.gov.br/api/pncp/v1"]
UA = {"User-Agent": "licitacoes-pncp-ia-worker/1.0 (+https://github.com/Max-Martins/licitacoes-pncp-ia)"}
TIMEOUT = httpx.Timeout(25.0, connect=10.0)
# "Disjuntor": se a API de documentos do PNCP estiver fora, não insiste em todos os processos.
_PNCP_FORA_ATE = 0.0
PAUSA_DISJUNTOR = 600  # segundos
MAX_BYTES = 150 * 1024 * 1024
EXT_TEXTO = {".pdf", ".docx", ".txt", ".odt", ".xlsx", ".csv", ".html", ".htm"}

# Pastas padronizadas dentro de cada processo
SUB_ORIGINAIS = "01_DOCUMENTOS_ORIGINAIS"
SUB_DECLARACOES = "02_MODELOS_DECLARACOES"
SUB_ANALISE = "03_ANALISE_IA"


class PNCPIndisponivel(RuntimeError):
    pass


def _get_pncp(client: httpx.Client, caminho: str) -> httpx.Response:
    """Consulta o PNCP tentando as duas rotas; abre o disjuntor se ambas falharem."""
    global _PNCP_FORA_ATE
    if time.time() < _PNCP_FORA_ATE:
        raise PNCPIndisponivel("API de documentos do PNCP indisponível (aguardando nova tentativa)")
    ultimo = None
    for base in PNCP_BASES:
        try:
            return _get(client, base + caminho, tentativas=2)
        except Exception as e:  # noqa: BLE001
            ultimo = e
    _PNCP_FORA_ATE = time.time() + PAUSA_DISJUNTOR
    raise PNCPIndisponivel(f"API de documentos do PNCP indisponível: {ultimo}")


def _get(client: httpx.Client, url: str, tentativas: int = 3) -> httpx.Response:
    espera = 2.0
    for i in range(tentativas):
        try:
            r = client.get(url)
            if r.status_code == 429 or r.status_code >= 500:
                raise httpx.HTTPStatusError(f"HTTP {r.status_code}", request=r.request, response=r)
            return r
        except (httpx.TransportError, httpx.HTTPStatusError) as e:
            if i == tentativas - 1:
                raise
            log.warning("retry %s em %.1fs (%s)", url, espera, e)
            time.sleep(espera)
            espera *= 2
    raise RuntimeError("inalcançável")


def listar_arquivos(client: httpx.Client, cnpj: str, ano: int, seq: int) -> list[dict]:
    r = _get_pncp(client, f"/orgaos/{cnpj}/compras/{ano}/{seq}/arquivos")
    if r.status_code == 204 or not r.content:
        return []
    r.raise_for_status()
    dados = r.json()
    return [a for a in dados if a.get("statusAtivo", True)]


def listar_itens(client: httpx.Client, cnpj: str, ano: int, seq: int, limite: int = 500) -> list[dict]:
    itens, pagina = [], 1
    while len(itens) < limite:
        r = _get_pncp(client, f"/orgaos/{cnpj}/compras/{ano}/{seq}/itens?pagina={pagina}&tamanhoPagina=100")
        if r.status_code == 204 or not r.content:
            break
        r.raise_for_status()
        lote = r.json()
        if not lote:
            break
        itens.extend(lote)
        if len(lote) < 100:
            break
        pagina += 1
    return [
        {
            "numero": i.get("numeroItem"),
            "descricao": i.get("descricao"),
            "quantidade": i.get("quantidade"),
            "unidade": i.get("unidadeMedida"),
            "valor_unitario_estimado": i.get("valorUnitarioEstimado"),
            "valor_total": i.get("valorTotal"),
            "beneficio_me_epp": i.get("tipoBeneficioNome"),
            "criterio": i.get("criterioJulgamentoNome"),
            "situacao": i.get("situacaoCompraItemNome"),
        }
        for i in itens
    ]


EXT_CONHECIDAS = {".pdf", ".zip", ".rar", ".7z", ".doc", ".docx", ".xls", ".xlsx", ".odt", ".ods", ".rtf",
                  ".txt", ".csv", ".html", ".htm", ".jpg", ".jpeg", ".png", ".tif", ".tiff", ".xml", ".p7s"}


def tipo_pelo_conteudo(conteudo: bytes, nome: str = "") -> str | None:
    """Descobre a extensão verdadeira pelos primeiros bytes (o nome/título do PNCP muitas vezes vem sem
    extensão ou com sufixos falsos como '.2026', o que deixa o arquivo ilegível no Windows)."""
    c = conteudo[:1024]
    if c.lstrip()[:5] == b"%PDF-":
        return ".pdf"
    if c[:4] == b"Rar!":
        return ".rar"
    if c[:6] == b"7z\xbc\xaf\x27\x1c":
        return ".7z"
    if c[:2] == b"PK":
        try:
            with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
                nomes = z.namelist()
        except zipfile.BadZipFile:
            return ".zip"
        if "word/document.xml" in nomes:
            return ".docx"
        if "xl/workbook.xml" in nomes:
            return ".xlsx"
        if "mimetype" in nomes and "content.xml" in nomes:
            try:
                with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
                    mt = z.read("mimetype").decode(errors="ignore")
            except Exception:  # noqa: BLE001
                mt = ""
            return ".ods" if "spreadsheet" in mt else ".odt"
        return ".zip"
    if c[:8] == bytes.fromhex("D0CF11E0A1B11AE1"):
        return ".xls" if nome.lower().endswith(".xls") else ".doc"
    if c[:5] == b"{\\rtf":
        return ".rtf"
    if c[:3] == b"\xff\xd8\xff":
        return ".jpg"
    if c[:8] == b"\x89PNG\r\n\x1a\n":
        return ".png"
    if c[:4] in (b"II*\x00", b"MM\x00*"):
        return ".tif"
    if c[:2] == b"\x1f\x8b":
        return ".gz"
    inicio = c.lstrip().lower()
    if inicio.startswith(b"<!doctype html") or inicio.startswith(b"<html"):
        return ".html"
    if b"%PDF-" in c:  # PDF com lixo antes do cabeçalho (aceito pelos leitores)
        return ".pdf"
    return None


def _nome_arquivo(meta: dict, resp: httpx.Response, ordem: int, conteudo: bytes = b"") -> str:
    cd = resp.headers.get("content-disposition", "")
    m = re.search(r"filename\*?=(?:UTF-8'')?\"?([^\";]+)", cd, re.I)
    original = m.group(1) if m else (meta.get("titulo") or f"documento_{ordem}")
    try:  # nomes em UTF-8 que chegam "%C3%A7..." ou com latin-1 trocado
        from urllib.parse import unquote
        original = unquote(original)
        original = original.encode("latin-1").decode("utf-8")
    except Exception:  # noqa: BLE001
        pass
    original = sem_acento(original)
    ext_nome = Path(original).suffix.lower()
    ext = tipo_pelo_conteudo(conteudo, original) if conteudo else None
    if not ext:
        if ext_nome in EXT_CONHECIDAS:
            ext = ext_nome
        else:
            ct = resp.headers.get("content-type", "")
            ext = ".pdf" if "pdf" in ct else ".zip" if "zip" in ct else ".rar" if "rar" in ct else (ext_nome if ext_nome in EXT_CONHECIDAS else ".pdf" if not ct else ".bin")
    base_nome = Path(original).stem if (ext_nome in EXT_CONHECIDAS or re.fullmatch(r"\.[a-z]{2,4}", ext_nome)) else original
    tipo = limpar_nome(meta.get("tipoDocumentoNome") or "DOCUMENTO", 30).replace(" ", "_")
    base = limpar_nome(base_nome, 70).replace(" ", "_").replace(".", "_")
    return f"{ordem:02d}_{tipo}__{base}{ext}"


def _sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def extrair_texto(caminho: Path) -> tuple[str, int, str]:
    """Retorna (texto, paginas, status). status: OK | SEM_TEXTO_OCR_NECESSARIO | NAO_SUPORTADO | ERRO:..."""
    ext = caminho.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(caminho))
            partes = []
            for n, pg in enumerate(reader.pages, 1):
                try:
                    t = pg.extract_text() or ""
                except Exception:  # página corrompida não derruba o arquivo
                    t = ""
                partes.append(f"\n[[{caminho.name} | página {n}]]\n{t}")
            texto = "".join(partes)
            uteis = len(re.sub(r"\s|\[\[.*?\]\]", "", texto))
            paginas = len(reader.pages)
            if paginas and uteis / paginas < 40:
                return texto, paginas, "SEM_TEXTO_OCR_NECESSARIO"
            return texto, paginas, "OK"
        if ext == ".docx":
            import docx
            d = docx.Document(str(caminho))
            linhas = [p.text for p in d.paragraphs]
            for tb in d.tables:
                for row in tb.rows:
                    linhas.append(" | ".join(c.text for c in row.cells))
            return f"\n[[{caminho.name}]]\n" + "\n".join(linhas), 0, "OK"
        if ext in {".txt", ".csv", ".html", ".htm"}:
            return f"\n[[{caminho.name}]]\n" + caminho.read_text(errors="replace"), 0, "OK"
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(caminho), read_only=True, data_only=True)
            linhas = []
            for ws in wb.worksheets:
                linhas.append(f"## {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        linhas.append(" | ".join("" if v is None else str(v) for v in row))
            return f"\n[[{caminho.name}]]\n" + "\n".join(linhas), 0, "OK"
        return "", 0, "NAO_SUPORTADO"
    except Exception as e:  # noqa: BLE001
        return "", 0, f"ERRO: {e.__class__.__name__}: {e}"[:200]


def _pasta_textos(destino: Path) -> Path:
    """Os .txt extraídos (só para a IA) ficam em 03_ANALISE_IA/textos_extraidos, não junto dos originais."""
    for anc in destino.parents:
        if anc.name == SUB_ORIGINAIS:
            p = anc.parent / SUB_ANALISE / "textos_extraidos"
            p.mkdir(parents=True, exist_ok=True)
            return p
    return destino.parent


def _registrar(destino: Path, conteudo: bytes, meta: dict, origem_url: str, dentro_de: str | None = None) -> dict:
    destino.write_bytes(conteudo)
    texto, paginas, status = ("", 0, "NAO_SUPORTADO")
    if destino.suffix.lower() in EXT_TEXTO:
        texto, paginas, status = extrair_texto(destino)
        if texto:
            (_pasta_textos(destino) / (destino.name + ".txt")).write_text(texto, encoding="utf-8")
    return {
        "arquivo": destino.name,
        "tipo_documento": meta.get("tipoDocumentoNome"),
        "titulo": meta.get("titulo"),
        "url_origem": origem_url,
        "extraido_de": dentro_de,
        "tamanho_bytes": len(conteudo),
        "sha256": _sha256(conteudo),
        "paginas": paginas,
        "leitura": status,
        "caracteres_texto": len(texto),
        "data_publicacao_pncp": meta.get("dataPublicacaoPncp"),
        "_texto": texto,
    }


def preparar_processo(raiz: Path, rec: dict, pasta_rel: str) -> dict:
    """Cria a pasta do processo, baixa e organiza todos os arquivos e devolve manifesto + texto consolidado."""
    pasta = raiz / pasta_rel
    originais = pasta / SUB_ORIGINAIS
    for p in (originais, pasta / SUB_DECLARACOES, pasta / SUB_ANALISE):
        p.mkdir(parents=True, exist_ok=True)

    cnpj, ano, seq = rec["cnpj"], int(rec["anoCompra"]), int(rec["sequencialCompra"])
    avisos: list[str] = []
    manifesto: list[dict] = []
    itens: list[dict] = []

    with httpx.Client(headers=UA, timeout=TIMEOUT, follow_redirects=True) as client:
        try:
            itens = listar_itens(client, cnpj, ano, seq)
        except Exception as e:  # noqa: BLE001
            avisos.append(f"Itens não obtidos: {e}")
        try:
            arquivos = listar_arquivos(client, cnpj, ano, seq)
        except Exception as e:  # noqa: BLE001
            arquivos = []
            avisos.append(f"Lista de arquivos não obtida: {e}")
        if not arquivos:
            avisos.append("PNCP não retornou arquivos para este processo — baixar no portal de origem.")

        for ordem, meta in enumerate(arquivos, 1):
            url = meta.get("url") or meta.get("uri")
            if not url:
                continue
            try:
                r = _get(client, url)
                r.raise_for_status()
                conteudo = r.content
                if len(conteudo) > MAX_BYTES:
                    avisos.append(f"{meta.get('titulo')}: acima de {MAX_BYTES // 2**20} MB, não baixado")
                    continue
                nome = _nome_arquivo(meta, r, ordem, conteudo)
                destino = originais / nome
                manifesto.append(_registrar(destino, conteudo, meta, url))
                manifesto.extend(_expandir_compactado(destino, conteudo, meta, url, avisos))
            except Exception as e:  # noqa: BLE001
                avisos.append(f"Falha ao baixar '{meta.get('titulo')}': {e}")

    for m in manifesto:
        if m["leitura"] == "SEM_TEXTO_OCR_NECESSARIO":
            avisos.append(f"{m['arquivo']}: PDF digitalizado (sem texto). A IA multimodal pode ler; conferir manualmente.")

    texto = "".join(m.pop("_texto") for m in manifesto)
    (pasta / "_MANIFESTO.json").write_text(json.dumps({
        "processo": rec.get("numeroControlePNCP"),
        "coletado_em": datetime.now().isoformat(timespec="seconds"),
        "arquivos": manifesto, "itens": itens, "avisos": avisos,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    (pasta / SUB_ANALISE / "itens_pncp.json").write_text(json.dumps(itens, ensure_ascii=False, indent=2), encoding="utf-8")

    return {"pasta": str(pasta), "arquivos": manifesto, "itens": itens, "avisos": avisos,
            "texto": texto, "caracteres": len(texto)}


def _nome_interno_zip(info: zipfile.ZipInfo) -> str:
    """Nomes dentro de ZIP gerado no Windows vêm em CP850/CP437 sem a flag UTF-8 (acentos viram lixo)."""
    nome = info.filename
    if not info.flag_bits & 0x800:
        try:
            nome = nome.encode("cp437").decode("cp850")
        except Exception:  # noqa: BLE001
            pass
    return nome


def _alvo_unico(pasta: Path, nome: str) -> Path:
    alvo = pasta / nome
    n = 2
    while alvo.exists():
        alvo = pasta / f"{Path(nome).stem}_{n}{Path(nome).suffix}"
        n += 1
    return alvo


def _salvar_extraido(pasta: Path, nome_original: str, dados: bytes, meta: dict, url: str, origem: str,
                     avisos: list[str], nivel: int) -> list[dict]:
    nome_original = sem_acento(Path(nome_original.replace("\\", "/")).name)
    if not nome_original or nome_original.startswith(".") or nome_original.lower() in {"thumbs.db", "desktop.ini"}:
        return []
    ext = tipo_pelo_conteudo(dados, nome_original) or Path(nome_original).suffix.lower() or ".bin"
    stem = Path(nome_original).stem if Path(nome_original).suffix.lower() in EXT_CONHECIDAS else nome_original
    # proteção contra zip-slip: nome sanitizado, sempre dentro da pasta do compactado
    alvo = _alvo_unico(pasta, limpar_nome(stem, 80).replace(" ", "_") + ext)
    saida = [_registrar(alvo, dados, meta, url, dentro_de=origem)]
    if nivel < 3:
        saida.extend(_expandir_compactado(alvo, dados, meta, url, avisos, nivel + 1))
    return saida


def _expandir_compactado(destino: Path, conteudo: bytes, meta: dict, url: str, avisos: list[str],
                         nivel: int = 1) -> list[dict]:
    """ZIP, RAR e 7z são extraídos para uma pasta com o mesmo nome (inclusive compactados dentro de compactados)."""
    ext = destino.suffix.lower()
    if ext not in {".zip", ".rar", ".7z"}:
        return []
    pasta = destino.with_suffix("")
    pasta.mkdir(exist_ok=True)
    saida: list[dict] = []
    if ext == ".zip":
        try:
            with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
                for info in z.infolist():
                    if info.is_dir() or info.file_size > MAX_BYTES:
                        continue
                    saida.extend(_salvar_extraido(pasta, _nome_interno_zip(info), z.read(info), meta, url,
                                                  destino.name, avisos, nivel))
        except (zipfile.BadZipFile, RuntimeError, NotImplementedError) as e:
            avisos.append(f"{destino.name}: ZIP inválido, corrompido ou com senha ({e.__class__.__name__}) — abrir manualmente")
        return saida
    # RAR / 7z: usa o 'unar' instalado na imagem do worker
    import shutil
    import subprocess
    import tempfile
    if not shutil.which("unar"):
        avisos.append(f"{destino.name}: {ext} não extraído (ferramenta unar ausente) — abrir manualmente")
        return saida
    with tempfile.TemporaryDirectory() as tmp:
        proc = subprocess.run(["unar", "-q", "-f", "-D", "-e", "UTF-8", "-o", tmp, str(destino)],
                              capture_output=True, timeout=300)
        if proc.returncode != 0:
            avisos.append(f"{destino.name}: falha ao extrair {ext} (pode ter senha) — abrir manualmente")
            return saida
        for f in sorted(Path(tmp).rglob("*")):
            if f.is_file() and f.stat().st_size <= MAX_BYTES:
                saida.extend(_salvar_extraido(pasta, f.name, f.read_bytes(), meta, url, destino.name, avisos, nivel))
    return saida


# ---------------------------------------------------------------------------------------------
# Documentos pendentes: processos cuja pasta foi criada com a API de documentos do PNCP fora do ar.
# Uma rotina do worker tenta de novo periodicamente e preenche 01_DOCUMENTOS_ORIGINAIS sozinha.
# ---------------------------------------------------------------------------------------------
_RE_CONTROLE = re.compile(r"^(\d{14})-\d+-(\d+)/(\d{4})$")


def _pendente(pasta: Path) -> dict | None:
    man = pasta / "_MANIFESTO.json"
    if not man.exists():
        return None
    try:
        dados = json.loads(man.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return None
    originais = pasta / SUB_ORIGINAIS
    tem_arquivos = originais.exists() and any(p.is_file() for p in originais.rglob("*"))
    if tem_arquivos:
        return None
    avisos = " ".join(dados.get("avisos") or [])
    if dados.get("arquivos") and tem_arquivos:
        return None
    if "indispon" not in avisos.lower() and dados.get("tentativas_docs", 0) >= 3:
        return None  # PNCP respondeu "sem arquivos" 3 vezes: não insiste
    m = _RE_CONTROLE.match(str(dados.get("processo") or ""))
    if not m:
        return None
    cnpj, seq, ano = m.group(1), int(m.group(2)), int(m.group(3))
    return {"cnpj": cnpj, "anoCompra": ano, "sequencialCompra": seq,
            "numeroControlePNCP": dados.get("processo"), "_tentativas": dados.get("tentativas_docs", 0)}


def completar_pendentes(raiz: Path, limite: int = 400, incluir_vencidos: bool = False) -> dict:
    """Varre LICITACOES/UF/DD.MM.AAAA/<processo> e baixa os documentos que faltam."""
    hoje = datetime.now().date()
    resumo = {"verificados": 0, "pendentes": 0, "baixados": 0, "sem_arquivos": 0, "pncp_fora": False, "pastas": []}
    for uf_dir in sorted(p for p in raiz.iterdir() if p.is_dir() and not p.name.startswith("_")):
        for dia in sorted(p for p in uf_dir.iterdir() if p.is_dir()):
            try:
                data = datetime.strptime(dia.name, "%d.%m.%Y").date()
            except ValueError:
                continue
            if data < hoje and not incluir_vencidos:
                continue  # sessão já passou
            for pasta in sorted(p for p in dia.iterdir() if p.is_dir()):
                resumo["verificados"] += 1
                rec = _pendente(pasta)
                if not rec:
                    continue
                resumo["pendentes"] += 1
                if resumo["pendentes"] > limite:
                    continue
                if time.time() < _PNCP_FORA_ATE:
                    resumo["pncp_fora"] = True
                    continue
                tentativas = rec.pop("_tentativas") + 1
                res = preparar_processo(raiz, rec, str(pasta.relative_to(raiz)))
                man = pasta / "_MANIFESTO.json"
                dados = json.loads(man.read_text(encoding="utf-8"))
                dados["tentativas_docs"] = tentativas
                man.write_text(json.dumps(dados, ensure_ascii=False, indent=2), encoding="utf-8")
                if res["arquivos"]:
                    resumo["baixados"] += 1
                    resumo["pastas"].append(str(pasta.relative_to(raiz)))
                    resumo_md = pasta / "00_RESUMO_DO_PROCESSO.md"
                    if resumo_md.exists():
                        with resumo_md.open("a", encoding="utf-8") as fh:
                            fh.write(f"\n\n> Documentos do PNCP baixados depois, em {datetime.now():%d/%m/%Y %H:%M} "
                                     f"({len(res['arquivos'])} arquivo(s) em {SUB_ORIGINAIS}). A análise de IA acima "
                                     "foi feita antes do edital estar disponível — conferir o edital.\n")
                elif any("indispon" in a.lower() for a in res["avisos"]):
                    resumo["pncp_fora"] = True
                else:
                    resumo["sem_arquivos"] += 1
    log.info("documentos pendentes: %s", {k: v for k, v in resumo.items() if k != "pastas"})
    return resumo
