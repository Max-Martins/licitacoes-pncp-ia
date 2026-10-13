"""Proposta prévia (DOCX) por empresa + documento de conferência do cruzamento item × catálogo.

Layout padrão de proposta comercial (Lei 14.133/2021). Para usar o papel timbrado da sua empresa, salve um .docx
em worker/app/modelos/papel_timbrado.docx (só cabeçalho/rodapé; o conteúdo é acrescentado pelo sistema).
Tudo que a IA não confirmou sai destacado em amarelo para conferência manual.
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SUB_PROPOSTA = "04_PROPOSTA_PREVIA"
COTAVEIS = ("ATENDE", "ATENDE_PARCIAL", "VERIFICAR")
COR = {"ATENDE": "E2F0D9", "ATENDE_PARCIAL": "FFF2CC", "VERIFICAR": "FCE4D6", "NAO_ATENDE": "F2F2F2"}
ROTULO = {"ATENDE": "Atende", "ATENDE_PARCIAL": "Atende parcialmente — conferir", "VERIFICAR": "Verificar",
          "NAO_ATENDE": "Não atende"}
MESES = "janeiro fevereiro março abril maio junho julho agosto setembro outubro novembro dezembro".split()
PLACEHOLDER = {"razao_social": "[RAZÃO SOCIAL]", "cnpj": "[CNPJ]", "ie": "[INSCRIÇÃO ESTADUAL]", "im": "[INSCRIÇÃO MUNICIPAL]",
               "endereco": "[ENDEREÇO COMPLETO]", "cidade_uf": "[CIDADE/UF]", "telefone": "[TELEFONE]", "email": "[E-MAIL]",
               "representante": "[REPRESENTANTE LEGAL]", "rg": "[RG]", "cpf": "[CPF]", "cargo": "[CARGO]",
               "banco": "[BANCO]", "agencia": "[AGÊNCIA]", "conta": "[CONTA CORRENTE]"}


# ----------------------------------------------------------------- utilidades
def brl(v) -> str:
    try:
        return f"{float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return ""


def _num(v):
    try:
        return float(str(v).replace(",", ".")) if v not in (None, "") else None
    except ValueError:
        return None


def _qtd(v) -> str:
    n = _num(v)
    if n is None:
        return str(v or "")
    return f"{n:,.0f}".replace(",", ".") if n.is_integer() else brl(n)


_U = "zero um dois três quatro cinco seis sete oito nove dez onze doze treze quatorze quinze dezesseis dezessete dezoito dezenove".split()
_D = "_ _ vinte trinta quarenta cinquenta sessenta setenta oitenta noventa".split()
_C = "_ cento duzentos trezentos quatrocentos quinhentos seiscentos setecentos oitocentos novecentos".split()


def _ate_mil(n: int) -> str:
    if n == 100:
        return "cem"
    c, r = divmod(n, 100)
    partes = [_C[c]] if c else []
    if r:
        partes.append(_U[r] if r < 20 else _D[r // 10] + (f" e {_U[r % 10]}" if r % 10 else ""))
    return " e ".join(partes)


def extenso(valor: float) -> str:
    """Valor em reais por extenso (até centenas de milhões)."""
    reais, cent = divmod(round(valor * 100), 100)
    grupos, n = [], reais
    while n:
        n, g = divmod(n, 1000)
        grupos.append(g)
    nomes = [("", ""), ("mil", "mil"), ("milhão", "milhões"), ("bilhão", "bilhões")]
    pedacos = []
    for i, g in reversed(list(enumerate(grupos))):
        if not g:
            continue
        txt = "mil" if (i == 1 and g == 1) else _ate_mil(g) + (f" {nomes[i][0] if g == 1 else nomes[i][1]}" if i else "")
        pedacos.append((txt, g, i))
    txt = ""
    for j, (p, g, i) in enumerate(pedacos):
        if j:
            txt += " e " if (i == 0 and (g < 100 or g % 100 == 0)) else ", "
        txt += p
    out = []
    if reais:
        sufixo = " de reais" if reais % 1_000_000 == 0 else (" real" if reais == 1 else " reais")
        out.append(txt + sufixo)
    if cent:
        out.append(_ate_mil(cent) + (" centavo" if cent == 1 else " centavos"))
    return " e ".join(out) or "zero real"


def _sombra(cell, cor: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), cor)
    tcPr.append(shd)


def _larguras(tab, larg) -> None:
    """Fixa a largura das colunas (grade da tabela + células) — o Word e o LibreOffice respeitam."""
    tab.autofit = False
    tblPr = tab._tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    for i, w in enumerate(larg):
        tab.columns[i].width = Cm(w)
        for cell in tab.columns[i].cells:
            cell.width = Cm(w)


def _repetir_cabecalho(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _txt(cell, texto, negrito=False, tam=9, alinhar=None, cor=None, destacar=False, italico=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if alinhar:
        p.alignment = alinhar
    run = p.add_run(str(texto if texto is not None else ""))
    run.bold, run.italic, run.font.size = negrito, italico, Pt(tam)
    if cor:
        run.font.color.rgb = RGBColor.from_string(cor)
    if destacar:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def _par(doc, texto="", negrito=False, tam=10, alinhar=None, depois=4, destacar=False, italico=False, cor=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(depois)
    if alinhar:
        p.alignment = alinhar
    if texto:
        r = p.add_run(texto)
        r.bold, r.italic, r.font.size = negrito, italico, Pt(tam)
        if destacar:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if cor:
            r.font.color.rgb = RGBColor.from_string(cor)
    return p


def _documento(paisagem=True) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    if paisagem:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Cm(1.8))
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    return doc


def _ref_processo(rec: dict) -> dict:
    return {
        "modalidade": f"{rec.get('modalidadeNome') or 'Licitação'} nº {rec.get('numeroCompra') or '[Nº]'}/{rec.get('anoCompra') or ''}",
        "processo": str(rec.get("processo") or "[Nº DO PROCESSO]"),
        "orgao": str(rec.get("orgao") or ""),
        "unidade": str(rec.get("unidade") or ""),
        "local": f"{rec.get('municipio') or ''}/{rec.get('uf') or ''}",
        "objeto": " ".join(str(rec.get("objetoCompra") or "").split()),
        "encerramento": rec.get("dataEncerramentoProposta"),
        "controle": rec.get("numeroControlePNCP") or "",
    }


# ----------------------------------------------------------------- proposta (papel timbrado opcional)
MODELO = Path(__file__).parent / "modelos" / "papel_timbrado.docx"
DIAS = "segunda-feira terça-feira quarta-feira quinta-feira sexta-feira sábado domingo".split()
LARGURA_UTIL = 17.4  # cm (A4 retrato com margens de 1,8 cm)


def _grade(doc, linhas, larg):
    """Tabela rótulo/valor do modelo. linhas: [[(texto, negrito, destacar, span), ...], ...]"""
    tab = doc.add_table(rows=0, cols=len(larg))
    tab.style = "Table Grid"
    tab.autofit = False
    for linha in linhas:
        row = tab.add_row()
        col = 0
        for texto, neg, dest, span in linha:
            cell = row.cells[col]
            if span > 1:
                cell = cell.merge(row.cells[col + span - 1])
            cell.width = Cm(sum(larg[col:col + span]))
            _txt(cell, texto, negrito=neg, tam=8.5, destacar=dest)
            col += span
    _larguras(tab, larg)
    return tab


def _v_edital(analise: dict, grupo: str, chave: str, padrao: str | None) -> tuple[str, bool]:
    v = ((analise.get(grupo) or {}).get(chave) or {})
    v = v.get("valor") if isinstance(v, dict) else v
    v = " ".join(str(v or "").split())
    if v and "NÃO HÁ INFORMAÇÃO" not in v.upper():
        return v.upper(), False
    return (padrao or "CONFORME EDITAL").upper(), True


def gerar_proposta(destino: Path, rec: dict, analise: dict, emp: dict, itens: list[dict]) -> dict:
    ref = _ref_processo(rec)
    doc = Document(str(MODELO)) if MODELO.exists() else _documento(paisagem=False)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    rep = emp.get("representante") if isinstance(emp.get("representante"), dict) else {}
    pad = emp.get("padroes") or {}

    _par(doc, "PRÉVIA GERADA AUTOMATICAMENTE — conferir itens, preços e campos em amarelo antes de enviar (apague esta linha).",
         tam=8, cor="C00000", negrito=True, alinhar=WD_ALIGN_PARAGRAPH.CENTER, depois=2)
    _par(doc, "PROPOSTA COMERCIAL", negrito=True, tam=12, alinhar=WD_ALIGN_PARAGRAPH.CENTER, depois=6)

    contato = analise.get("orgao_contato") or {}
    def campo(rotulo, valor, sublinhar=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(rotulo)
        r.bold, r.font.size = True, Pt(10)
        v = " ".join(str(valor or "").split())
        falta = not v or "NÃO HÁ INFORMAÇÃO" in v.upper()
        r = p.add_run(f"[{rotulo.strip(': ').upper()} — conferir no edital]" if falta else v)
        r.font.size, r.underline = Pt(10), sublinhar and not falta
        if falta:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    _par(doc, "À", depois=0)
    req = ref["orgao"] + (f" – {ref['unidade']}" if ref["unidade"] and ref["unidade"].upper() not in ref["orgao"].upper() else "")
    campo("Requisitante: ", req.upper(), sublinhar=True)
    campo("Endereço: ", contato.get("endereco"))
    campo("Telefone: ", contato.get("telefone"))
    campo("Site: ", contato.get("site") or contato.get("email"))
    _par(doc, depois=4)

    modalidade = (rec.get("modalidadeNome") or "Licitação").replace(" - ", " ").upper()
    numero = f"{rec.get('numeroCompra') or '[Nº]'}/{rec.get('anoCompra') or ''}"
    srp = " – SISTEMA DE REGISTRO DE PREÇOS" if rec.get("srp") else ""
    caixa = doc.add_table(rows=1, cols=1)
    caixa.style = "Table Grid"
    c = caixa.rows[0].cells[0]
    c.width = Cm(LARGURA_UTIL)
    _sombra(c, "D9D9D9")
    _txt(c, "PROPOSTA DE PREÇOS", negrito=True, tam=11, alinhar=WD_ALIGN_PARAGRAPH.CENTER)
    for t, tam in ((f"{modalidade} Nº {numero}", 11), (f"PROCESSO ADMINISTRATIVO Nº {ref['processo']}{srp}", 8.5)):
        p = c.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t)
        r.bold, r.font.size = True, Pt(tam)
    row = caixa.add_row()
    _larguras(caixa, [LARGURA_UTIL])
    p = _txt(row.cells[0], "", tam=9.5, alinhar=WD_ALIGN_PARAGRAPH.JUSTIFY)
    for t, neg in (("Em atendimento ao ", False), (f"Edital de Licitação – {modalidade.title()} nº {numero}", True),
                   (f" (Processo Administrativo nº {ref['processo']}{srp.title().replace('De ', 'de ')}) da {ref['orgao'].title()}, "
                    "apresentamos nossa proposta de preços para os itens abaixo especificados, conforme o Termo de Referência. "
                    "Nos preços estão incluídas todas as despesas (preços líquidos).", False)):
        r = p.add_run(t)
        r.bold, r.font.size = neg, Pt(9.5)

    def d(chave, padrao=""):
        v = str(emp.get(chave) or "").strip()
        return (v.upper(), False) if v else ((padrao or PLACEHOLDER.get(chave, f"[{chave.upper()}]")), True)
    garantia, fg = _v_edital(analise, "garantias_e_custos", "garantia_contratual", pad.get("garantia") or " ")
    prazo, fp = _v_edital(analise, "prazos", "prazo_entrega", pad.get("prazo_entrega"))
    validade, fv = _v_edital(analise, "prazos", "validade_proposta", pad.get("validade_proposta"))
    pagto, fpg = _v_edital(analise, "prazos", "prazo_pagamento", pad.get("condicao_pagamento"))
    if garantia.strip() == "" or garantia.upper().startswith("NÃO"):
        garantia, fg = "", False
    L = [3.9, 5.6, 3.9, 4.6]
    v = {k: d(k) for k in ("razao_social", "nome_fantasia", "cnpj", "vendedor", "im", "ie", "endereco", "bairro",
                            "municipio_uf", "cep", "telefone", "email")}
    _grade(doc, [
        [("RAZÃO SOCIAL:", True, False, 1), (v["razao_social"][0], False, v["razao_social"][1], 3)],
        [("NOME FANTASIA:", True, False, 1), ((v["nome_fantasia"] if not v["nome_fantasia"][1] else v["razao_social"])[0], False, False, 3)],
        [("CNPJ:", True, False, 1), (v["cnpj"][0], False, v["cnpj"][1], 1), ("NOME DO VENDEDOR:", True, False, 1), (v["vendedor"][0], False, v["vendedor"][1], 1)],
        [("INSCRIÇÃO MUNICIPAL:", True, False, 1), (v["im"][0], False, v["im"][1], 1), ("INSCRIÇÃO ESTADUAL:", True, False, 1), (v["ie"][0], False, v["ie"][1], 1)],
        [("ENDEREÇO:", True, False, 1), (v["endereco"][0], False, v["endereco"][1], 1), ("BAIRRO:", True, False, 1), (v["bairro"][0], False, v["bairro"][1], 1)],
        [("MUNICÍPIO:", True, False, 1), (v["municipio_uf"][0], False, v["municipio_uf"][1], 1), ("CEP:", True, False, 1), (v["cep"][0], False, v["cep"][1], 1)],
        [("TELEFONE:", True, False, 1), (v["telefone"][0], False, v["telefone"][1], 1), ("E-MAIL:", True, False, 1), (v["email"][0], False, v["email"][1], 1)],
        [("GARANTIA (SE FOR O CASO):", True, False, 1), (garantia, False, fg, 1), ("PRAZO DA ENTREGA (CIF):", True, False, 1), (prazo, False, fp, 1)],
        [("VALIDADE DA PROPOSTA", True, False, 1), (validade, False, fv, 1), ("CONDIÇÃO DE PAGAMENTO", True, False, 1), (pagto, True, fpg, 1)],
    ], L)
    _par(doc, depois=4)

    cab = ["Lote", "Item", "Descrição", "Unid.", "Quant.", "Valor Unitário", "Total", "Marca/Obs."]
    larg = [1.1, 1.1, 6.6, 1.3, 1.6, 2.1, 2.3, 1.9]
    tab = doc.add_table(rows=1, cols=len(cab))
    tab.style = "Table Grid"
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    for i, (t, w) in enumerate(zip(cab, larg)):
        cell = tab.rows[0].cells[i]
        cell.width = Cm(w)
        _txt(cell, t, negrito=True, tam=9, alinhar=WD_ALIGN_PARAGRAPH.CENTER)
    _repetir_cabecalho(tab.rows[0])
    total, faltam = 0.0, 0
    if not itens:
        row = tab.add_row()
        m = row.cells[0].merge(row.cells[-1])
        _txt(m, "[Itens do edital não foram obtidos automaticamente — preencher manualmente]", italico=True, tam=9, destacar=True)
    C = WD_ALIGN_PARAGRAPH.CENTER
    for it in itens:
        o = it["oferta"]
        qtd = _num(it.get("quantidade"))
        unit = round(o["preco_unitario_sugerido"], 2) if o.get("preco_unitario_sugerido") else None
        vt = round(unit * qtd, 2) if (unit and qtd) else None
        total += vt or 0
        faltam += 0 if vt else 1
        row = tab.add_row()
        obs = f"{o.get('marca') or ''}"
        vals = [it.get("lote") or "", it.get("numero"), (o.get("descricao") or "").upper(), (it.get("unidade") or "").upper(),
                _qtd(it.get("quantidade")), f"R$ {brl(unit)}" if unit else "[PREÇO]", f"R$ {brl(vt)}" if vt else "[TOTAL]", obs]
        alin = [C, C, WD_ALIGN_PARAGRAPH.JUSTIFY, C, C, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, C]
        for i, (val, w) in enumerate(zip(vals, larg)):
            cell = row.cells[i]
            cell.width = Cm(w)
            p = _txt(cell, val, tam=8.5, alinhar=alin[i], destacar=(i in (5, 6) and not unit))
            if o["veredito"] != "ATENDE":
                _sombra(cell, COR.get(o["veredito"], "FFF2CC"))
        rr = row.cells[7].paragraphs[0].add_run(f"\nCód. {o.get('codigo')}" + (f" ({o['tabela_nome']})" if o.get("tabela_nome") else ""))
        rr.font.size, rr.italic = Pt(7), True
        rr.font.highlight_color = WD_COLOR_INDEX.YELLOW
    _larguras(tab, larg)
    row = tab.add_row()
    m = row.cells[0].merge(row.cells[5])
    _txt(m, "VALOR TOTAL DA PROPOSTA", negrito=True, tam=9, alinhar=WD_ALIGN_PARAGRAPH.RIGHT)
    m2 = row.cells[6].merge(row.cells[7])
    _txt(m2, f"R$ {brl(total)}" if total else "[TOTAL]", negrito=True, tam=9, alinhar=WD_ALIGN_PARAGRAPH.CENTER, destacar=not total)
    _par(doc, depois=6)
    if total:
        p = _par(doc, f"VALOR TOTAL DA PROPOSTA: R$ {brl(total)} ({extenso(total).upper()}).", negrito=True, tam=9.5,
                 depois=10, alinhar=WD_ALIGN_PARAGRAPH.JUSTIFY, destacar=bool(faltam))

    hoje = datetime.now()
    cidade = (emp.get("cidade_assinatura") or "[CIDADE]")
    _par(doc, f"{cidade}, {DIAS[hoje.weekday()]}, {hoje.day} de {MESES[hoje.month - 1]} de {hoje.year}",
         tam=9.5, alinhar=WD_ALIGN_PARAGRAPH.CENTER, depois=40)
    _par(doc, "Dados e assinatura do responsável pela empresa:", negrito=True, tam=9, depois=0)
    linhas = [rep.get("nome") or "[REPRESENTANTE LEGAL]"]
    if rep.get("nascimento"):
        linhas.append(f"Data de nascimento: {rep['nascimento']}")
    linhas.append(f"RG: {rep.get('rg') or '[RG]'} / CPF: {rep.get('cpf') or '[CPF]'}")
    if rep.get("endereco"):
        linhas.append(f"Endereço: {rep['endereco']}")
    linhas.append(f"Cargo: {rep.get('cargo') or '[CARGO]'}" + (f" / Estado Civil: {rep['estado_civil']}" if rep.get("estado_civil") else ""))
    for t in linhas:
        _par(doc, t, tam=9, depois=0)
    if rep.get("emails"):
        _par(doc, depois=0)
        _par(doc, f"E-mail: {rep['emails']}", tam=9, depois=0)
    if emp.get("rubrica"):
        _par(doc, emp["rubrica"], tam=7, alinhar=WD_ALIGN_PARAGRAPH.RIGHT, depois=0, cor="2E74B5")

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)
    return {"arquivo": destino.name, "itens": len(itens), "total": round(total, 2), "sem_preco": faltam}


# ----------------------------------------------------------------- conferência
def gerar_conferencia(destino: Path, rec: dict, cruz: dict) -> None:
    ref = _ref_processo(rec)
    doc = _documento()
    _par(doc, "CONFERÊNCIA DO CRUZAMENTO — ITENS DO EDITAL × TABELAS DE PREÇO", negrito=True, tam=13, depois=2, cor="1F3864")
    _par(doc, f"{ref['orgao']} — {ref['modalidade']} — Processo {ref['processo']}", negrito=True, depois=0)
    _par(doc, ref["objeto"], tam=9, depois=4)
    cats = cruz.get("catalogos") or {}
    nomes = {k: (v or {}).get("nome") or k for k, v in cats.items()}
    _par(doc, "Tabelas usadas: " + ("; ".join(f"{nomes[k]}: {v.get('arquivo')} ({v.get('produtos')} produtos)"
                                                for k, v in cats.items()) or "nenhuma encontrada em LICITACOES/_CATALOGOS"),
         tam=9, depois=2)
    _par(doc, ("Decisão por agente de IA (Claude) sobre candidatos pré-selecionados por similaridade. "
               if cruz.get("por_ia") else "Agente de IA indisponível: sugestões apenas por similaridade de texto. ")
         + "Nada aqui substitui a leitura do edital e do TR: confira especificação, apresentação e preço antes de enviar.",
         tam=9, depois=2, italico=True)
    for a in cruz.get("avisos") or []:
        _par(doc, "⚠ " + a, tam=9, depois=1, destacar=True)
    for o in cruz.get("observacoes") or []:
        _par(doc, "• " + str(o), tam=9, depois=1)
    _par(doc, depois=4)

    tabelas = list(cats) or ["—"]
    cab = ["Item", "Edital (especificação · qtd · ref. R$)", *[nomes.get(t, t) for t in tabelas], "Recomendação"]
    larg = [1.1, 8.3, *[round(13.2 / len(tabelas), 2)] * len(tabelas), 3.5]
    tab = doc.add_table(rows=1, cols=len(cab))
    tab.style = "Table Grid"
    tab.autofit = False
    for i, (t, w) in enumerate(zip(cab, larg)):
        c = tab.rows[0].cells[i]
        c.width = Cm(w)
        _sombra(c, "1F3864")
        _txt(c, t, negrito=True, tam=8.5, cor="FFFFFF", alinhar=WD_ALIGN_PARAGRAPH.CENTER)
    _repetir_cabecalho(tab.rows[0])
    for it in cruz.get("itens") or []:
        row = tab.add_row()
        for i, w in enumerate(larg):
            row.cells[i].width = Cm(w)
        _txt(row.cells[0], it.get("numero"), negrito=True, tam=8.5, alinhar=WD_ALIGN_PARAGRAPH.CENTER)
        p = _txt(row.cells[1], it.get("descricao"), tam=8)
        extra = f"Qtd: {_qtd(it.get('quantidade'))} {it.get('unidade') or ''}"
        if it.get("valor_unitario_estimado"):
            extra += f" · Ref. unit.: R$ {brl(it['valor_unitario_estimado'])}"
        if it.get("marca_referencia"):
            extra += f" · Marca ref.: {it['marca_referencia']}"
        r = p.add_run("\n" + extra)
        r.font.size, r.bold = Pt(7.5), True
        for col, emp in enumerate(tabelas, start=2):
            o = (it.get("ofertas") or {}).get(emp)
            cell = row.cells[col]
            if not o:
                _txt(cell, "Sem candidato na tabela", italico=True, tam=8, cor="808080")
                continue
            _sombra(cell, COR.get(o.get("veredito"), "FFF2CC"))
            p = _txt(cell, ROTULO.get(o.get("veredito"), o.get("veredito")), negrito=True, tam=8)
            linhas = []
            if o.get("codigo"):
                linhas.append(f"Cód. {o['codigo']} · {o.get('descricao')} · {o.get('marca')}")
                preco = f"Tabela R$ {brl(o['preco_tabela'])}" if o.get("preco_tabela") else "SEM PREÇO NA TABELA"
                if o.get("converter_para_unidade"):
                    preco += f" ÷ {o.get('itens_por_embalagem')} = R$ {brl(o.get('preco_unitario_sugerido'))}/un"
                ref_u, nosso = _num(it.get("valor_unitario_estimado")), o.get("preco_unitario_sugerido")
                if ref_u and nosso and nosso > ref_u:
                    preco += f"  ⚠ ACIMA da referência do edital (R$ {brl(ref_u)})"
                linhas.append(preco)
            if o.get("justificativa"):
                linhas.append(o["justificativa"])
            linhas += [f"– {d}" for d in o.get("divergencias") or []]
            if o.get("alternativas"):
                linhas.append("Alternativas: " + " | ".join(o["alternativas"]))
            for ln in linhas:
                rr = p.add_run("\n" + ln)
                rr.font.size = Pt(7.5)
        esc = it.get("escolha")
        o = (it.get("ofertas") or {}).get(esc) if esc else None
        p = _txt(row.cells[-1], (f"NA PROPOSTA: cód. {o['codigo']} ({nomes.get(esc, esc)})" if o and o.get("codigo")
                                else "Não cotado"), negrito=True, tam=8)
        if it.get("recomendacao"):
            rr = p.add_run("\n" + it["recomendacao"])
            rr.font.size = Pt(7.5)
    _larguras(tab, larg)
    if not cruz.get("itens"):
        row = tab.add_row()
        m = row.cells[0].merge(row.cells[-1])
        _txt(m, "Itens do edital não obtidos — cruzamento não realizado.", italico=True, tam=9)
    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destino)


def gerar_propostas(pasta: Path, rec: dict, analise: dict, empresa: dict) -> dict:
    """Uma única proposta em nome da empresa da Configuração, com o produto escolhido em qualquer das tabelas."""
    cruz = analise.get("cruzamento")
    if cruz is None:
        return {}
    destino = pasta / SUB_PROPOSTA
    gerar_conferencia(destino / "CONFERENCIA_CRUZAMENTO.docx", rec, cruz)
    sel = []
    for it in cruz.get("itens") or []:
        emp_id = it.get("escolha")
        o = (it.get("ofertas") or {}).get(emp_id) if emp_id else None
        if o and o.get("codigo") and o.get("veredito") in COTAVEIS:
            nome = ((cruz.get("catalogos") or {}).get(emp_id) or {}).get("nome")
            sel.append({**it, "oferta": {**o, "tabela": emp_id, "tabela_nome": nome if len(cruz.get("catalogos") or {}) > 1 else None}})
    info = gerar_proposta(destino / "PROPOSTA_PREVIA.docx", rec, analise, empresa or {}, sel)
    return {"conferencia": "CONFERENCIA_CRUZAMENTO.docx", "propostas": [info]}
