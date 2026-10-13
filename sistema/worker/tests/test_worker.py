"""Testes do worker (sem internet). Uso, na pasta sistema/worker:
    python -m unittest discover -s tests -v
Os testes da API HTTP rodam quando o FastAPI está instalado (pip install -r requirements.txt).
"""
import os
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.analise import salvar_analise  # noqa: E402
from app.catalogo import candidatos, tabela_do_arquivo  # noqa: E402
from app.naming import pasta_processo  # noqa: E402
from app.proposta import extenso, gerar_propostas  # noqa: E402
from app.relatorio import gerar_relatorio  # noqa: E402

REC = {
    "numeroControlePNCP": "11111111000101-1-00001/2026", "uf": "SP", "municipio": "Cidade Exemplo",
    "orgao": "MUNICIPIO DE EXEMPLO", "unidade": "Secretaria de Saúde", "anoCompra": 2026, "sequencialCompra": 1,
    "numeroCompra": "90012", "processo": "123/2026", "modalidadeId": 6, "modalidadeNome": "Pregão - Eletrônico",
    "objetoCompra": "Aquisição de reagentes para glicose", "dataEncerramentoProposta": "2026-10-20T09:30:00",
    "dataAberturaProposta": "2026-10-06T08:00:00", "srp": True, "valorTotalEstimado": 1000.0,
    "linkSistemaOrigem": "https://cnetmobile.estaleiro.serpro.gov.br/x", "usuarioNome": "Compras.gov.br",
    "link_pncp": "https://pncp.gov.br/app/editais/11111111000101/2026/1",
}
EMPRESA = {"nome": "Empresa Teste", "razao_social": "EMPRESA TESTE LTDA", "cnpj": "00.000.000/0001-00",
           "cidade_assinatura": "Cidade Teste", "representante": {"nome": "Fulano de Tal", "cargo": "Sócio"}}


def _planilha(arq: Path, linhas: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["Relatório exportado do ERP"])  # linha de título antes do cabeçalho, como nas exportações reais
    ws.append(["ITEM", "DESCRIÇÃO ITEM", "MARCA", "PREÇO"])
    for ln in linhas:
        ws.append(ln)
    wb.save(arq)


def _texto_docx(arq: Path) -> str:
    with zipfile.ZipFile(arq) as z:
        return z.read("word/document.xml").decode("utf-8")


class Catalogo(unittest.TestCase):
    def test_nome_do_arquivo_vira_tabela(self):
        self.assertEqual(tabela_do_arquivo("Tabela Principal"), "TABELA_PRINCIPAL")
        self.assertEqual(tabela_do_arquivo("Distribuidor X 2026-10-01"), "DISTRIBUIDOR_X")
        self.assertEqual(tabela_do_arquivo("preços_v2"), "PRECOS")

    def test_candidatos_em_varias_tabelas(self):
        with tempfile.TemporaryDirectory() as d:
            raiz = Path(d)
            (raiz / "_CATALOGOS").mkdir()
            _planilha(raiz / "_CATALOGOS" / "Tabela Principal.xlsx", [
                [101, "GLICOSE ENZ 1X500ML", "MARCA A", 80.5],
                [102, "COLESTEROL 1X100ML", "MARCA A", 60],
                [103, "TUBO COLETA VACUO EDTA 4ML C/100", "MARCA B", 120],
            ])
            _planilha(raiz / "_CATALOGOS" / "Distribuidor X.xlsx", [["X9", "GLICOSE LIQUIDA 2X250ML", "MARCA C", 70]])
            r = candidatos(raiz, [{"numero": "1", "descricao": "Reagente para glicose enzimática, frasco 500 mL", "quantidade": 10}], [])
            self.assertEqual(set(r["catalogos"]), {"TABELA_PRINCIPAL", "DISTRIBUIDOR_X"})
            self.assertEqual(r["catalogos"]["TABELA_PRINCIPAL"]["nome"], "Tabela Principal")
            principal = r["itens"][0]["candidatos"]["TABELA_PRINCIPAL"]
            self.assertEqual(principal[0]["codigo"], "101")
            self.assertEqual(r["itens"][0]["candidatos"]["DISTRIBUIDOR_X"][0]["codigo"], "X9")

    def test_sem_tabelas_nao_quebra(self):
        with tempfile.TemporaryDirectory() as d:
            r = candidatos(Path(d), [], [{"numero": "1", "descricao": "Luva"}])
            self.assertEqual(r["catalogos"], {})
            self.assertEqual(r["itens"][0]["candidatos"], {})


class Documentos(unittest.TestCase):
    def test_pasta_do_processo(self):
        info = pasta_processo(REC)
        self.assertEqual(info["portal"], "COMPRAS.GOV.BR")
        self.assertEqual(info["numero"], "PE 90012-2026")
        self.assertTrue(info["relativo"].startswith("SP/20.10.2026/20.10.2026-MUNICIPIO DE EXEMPLO_COMPRAS.GOV.BR"))

    def test_valor_por_extenso(self):
        self.assertEqual(extenso(1250.5), "mil, duzentos e cinquenta reais e cinquenta centavos")
        self.assertEqual(extenso(2_000_000), "dois milhões de reais")

    def test_analise_proposta_e_conferencia(self):
        analise = {
            "resumo": "Compra de reagentes.", "relevancia": "ALTA",
            "prazos": {"validade_proposta": {"valor": "90 dias", "fonte": "edital.pdf · pág. 2"}},
            "declaracoes": [{"nome": "Declaração de não emprego de menor", "conteudo": "que não emprega menor",
                             "momento": "habilitação", "fonte": "edital.pdf · pág. 9"}],
            "cruzamento": {
                "catalogos": {"TABELA_PRINCIPAL": {"arquivo": "Tabela Principal.xlsx", "produtos": 3, "nome": "Tabela Principal"},
                              "DISTRIBUIDOR_X": {"arquivo": "Distribuidor X.xlsx", "produtos": 1, "nome": "Distribuidor X"}},
                "por_ia": True, "avisos": [], "observacoes": [],
                "itens": [{"numero": "1", "descricao": "Reagente glicose 500 mL", "quantidade": 10, "unidade": "FR",
                           "valor_unitario_estimado": 90, "escolha": "DISTRIBUIDOR_X",
                           "ofertas": {"TABELA_PRINCIPAL": {"codigo": "101", "descricao": "GLICOSE ENZ 1X500ML", "marca": "MARCA A",
                                                            "preco_tabela": 80.5, "preco_unitario_sugerido": 80.5, "veredito": "ATENDE"},
                                       "DISTRIBUIDOR_X": {"codigo": "X9", "descricao": "GLICOSE LIQUIDA 2X250ML", "marca": "MARCA C",
                                                          "preco_tabela": 70, "preco_unitario_sugerido": 70, "veredito": "ATENDE_PARCIAL"}}}],
            },
        }
        with tempfile.TemporaryDirectory() as d:
            pasta = Path(d) / "processo"
            pasta.mkdir()
            out = salvar_analise(pasta, REC, analise, {"razao": "EMPRESA TESTE LTDA", "cnpj": "00.000.000/0001-00"})
            self.assertTrue((pasta / "00_RESUMO_DO_PROCESSO.md").exists())
            self.assertEqual(len(out["declaracoes"]), 1)
            md = (pasta / "00_RESUMO_DO_PROCESSO.md").read_text(encoding="utf-8")
            self.assertIn("**Tabela Principal:**", md)
            self.assertIn("**Distribuidor X:**", md)
            self.assertIn("EMPRESA TESTE LTDA", _texto_docx(pasta / "02_MODELOS_DECLARACOES" / out["declaracoes"][0]))

            res = gerar_propostas(pasta, REC, analise, EMPRESA)
            proposta = _texto_docx(pasta / "04_PROPOSTA_PREVIA" / "PROPOSTA_PREVIA.docx")
            self.assertIn("EMPRESA TESTE LTDA", proposta)
            self.assertIn("X9", proposta)
            self.assertIn("Distribuidor X", proposta)
            self.assertIn("700,00", proposta)  # 10 FR x R$ 70,00
            self.assertIn("Cidade Teste", proposta)
            self.assertEqual(res["propostas"][0]["total"], 700.0)
            conferencia = _texto_docx(pasta / "04_PROPOSTA_PREVIA" / "CONFERENCIA_CRUZAMENTO.docx")
            self.assertIn("Tabela Principal", conferencia)
            self.assertIn("Distribuidor X", conferencia)


class Relatorio(unittest.TestCase):
    def test_excel_abre_e_tem_as_abas(self):
        op = {"encerramento": datetime(2026, 10, 20, 9, 30), "abertura": None, "numero_processo": "PE 90012-2026",
              "processo_adm": "123/2026", "uf": "SP", "municipio": "Cidade Exemplo", "orgao": "MUNICIPIO DE EXEMPLO",
              "unidade": "Saúde", "modalidade": "Pregão - Eletrônico", "modo_disputa": "Aberto", "srp": True,
              "portal": "COMPRAS.GOV.BR", "valor": 1000.0, "objeto": "Aquisição de reagentes", "termos": ["reagente"],
              "situacao_pncp": "Divulgada no PNCP", "link_pncp": REC["link_pncp"], "link_portal": None,
              "pasta": "LICITACOES/SP/x", "controle_pncp": REC["numeroControlePNCP"], "declaracoes": [], "relevancia": "ALTA"}
        with tempfile.TemporaryDirectory() as d:
            destino = Path(d) / "rel.xlsx"
            gerar_relatorio(destino, [op], datetime(2026, 10, 14), datetime(2026, 10, 28),
                            {"ufs": ["SP"], "empresa": "Empresa Teste", "cnpj": "00.000.000/0001-00", "gerado_em": "13/10/2026 10:00"})
            wb = load_workbook(destino)
            self.assertEqual(wb.sheetnames[:5], ["Dashboard", "Oportunidades", "Detalhes", "Declarações", "Parâmetros"])
            textos = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value)
            self.assertIn("Empresa Teste", textos)
            self.assertIn("MUNICIPIO DE EXEMPLO", textos)


try:
    from fastapi.testclient import TestClient
except ImportError:  # pragma: no cover
    TestClient = None


@unittest.skipIf(TestClient is None, "FastAPI não instalado (pip install -r requirements.txt)")
class Api(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        os.environ["LICITACOES_ROOT"] = self.tmp.name
        os.environ["DOCS_RETRY_MIN"] = "0"
        import importlib
        import app.main
        self.main = importlib.reload(app.main)
        self.cli = TestClient(self.main.app)

    def tearDown(self):
        self.tmp.cleanup()

    def test_health_e_relatorio(self):
        r = self.cli.get("/health").json()
        self.assertTrue(r["ok"] and r["gravavel"])
        for sub in ("_RELATORIOS", "_SISTEMA", "_CATALOGOS"):
            self.assertTrue((Path(self.tmp.name) / sub).is_dir())
        rel = self.cli.post("/relatorio", json={
            "janela_ini": "2026-10-14T00:00:00", "janela_fim": "2026-10-28T23:59:59", "ufs": ["SP", "MG"],
            "empresa": "Empresa Teste", "oportunidades": [{**REC, "analise": {"relevancia": "ALTA"}}]}).json()
        self.assertEqual(rel["qtd"], 1)
        arq = self.cli.get("/relatorio/arquivo", params={"nome": rel["nome"]})
        self.assertEqual(arq.status_code, 200)
        self.assertEqual(self.cli.get("/relatorio/arquivo", params={"nome": "../../etc/passwd"}).status_code, 400)

    def test_candidatos_sem_tabelas(self):
        r = self.cli.post("/processo/candidatos", json={"itens_edital": [{"numero": "1", "descricao": "Luva"}]}).json()
        self.assertEqual(r["total_itens"], 1)


if __name__ == "__main__":
    unittest.main()
